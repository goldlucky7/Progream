# -*- coding: utf-8 -*-
"""
일본 영상 프로젝트 검토 프로그램

일본어 유튜브 영상 제작물(대본, 이미지, 자막, 씬 매핑, 타임라인, 오디오)을
한 폴더에 모아두면, 업로드 전에 문제가 없는지 자동으로 검토해서
보기 좋은 HTML 리포트(검토리포트.html)를 만들어 줍니다.

사용법:
    python japan_review.py "프로젝트폴더경로"
    python japan_review.py "프로젝트폴더경로" --cpm 320 --no-open
"""

import argparse
import csv
import os
import re
import sys
import html as html_mod
import unicodedata
from datetime import datetime
from pathlib import Path

# 선택 라이브러리 - 없으면 해당 검사만 건너뜁니다
try:
    from PIL import Image
except ImportError:
    Image = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import mutagen
except ImportError:
    mutagen = None

# Windows 콘솔 한글 깨짐 방지
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

# ── 기준값 (필요하면 여기 숫자만 바꾸면 됩니다) ──────────────────────────────
DEFAULT_CPM = 320          # 일본어 TTS 분당 글자수(공백 제외) 기준
CPM_RANGE = (280, 360)     # 예상 길이 표시용 범위
SENTENCE_MAX = 70          # 이 글자수를 넘는 문장은 "너무 긴 문장"으로 표시
SUB_LINE_MAX = 25          # 자막 한 줄 최대 글자수 (일본어 권장 20자 이내)
SUB_CPS_WARN = 9.0         # 자막 초당 글자수 - 이보다 빠르면 읽기 힘듦
SUB_CPS_BAD = 13.0         # 이보다 빠르면 사실상 못 읽음
SUB_MIN_DUR = 0.7          # 자막 최소 표시 시간(초)
SUB_MAX_DUR = 10.0         # 자막 최대 표시 시간(초)
IMG_MIN_WIDTH = 1280       # 이미지 최소 가로 해상도
IMG_RECOMMEND = (1920, 1080)
SCENE_NO_SUB_MIN = 4.0     # 이 길이(초) 이상인 씬에 자막이 하나도 없으면 주의
AUDIO_DIFF_RATIO = 0.20    # 대본 예상 길이와 오디오 길이 차이 허용 비율

HANGUL_RE = re.compile(r"[가-힣ㄱ-ㅎㅏ-ㅣ]")
JAPANESE_RE = re.compile(r"[ぁ-んァ-ヶ一-龯]")
HALFWIDTH_KANA_RE = re.compile(r"[\uFF66-\uFF9D]")
EMOJI_RE = re.compile(
    "[\U0001F300-\U0001FAFF\U00002600-\U000027BF\U0001F000-\U0001F0FF\u2b50\u2705\u274c]"
)

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
AUDIO_EXTS = {".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg"}
VIDEO_EXTS = {".mp4", ".mov", ".mkv", ".avi", ".webm"}


# ── 검사 결과 담는 그릇 ──────────────────────────────────────────────────────
class Report:
    SEVERITIES = ["오류", "주의", "정보", "통과"]

    def __init__(self):
        self.findings = []          # (섹션, 심각도, 제목, 상세)
        self.sections = []          # 섹션 표시 순서
        self.stats = {}             # 섹션별 요약 수치 {섹션: [(이름, 값), ...]}

    def add(self, section, severity, title, detail=""):
        if section not in self.sections:
            self.sections.append(section)
        self.findings.append((section, severity, title, detail))

    def stat(self, section, name, value):
        if section not in self.sections:
            self.sections.append(section)
        self.stats.setdefault(section, []).append((name, value))

    def count(self, severity):
        return sum(1 for f in self.findings if f[1] == severity)


# ── 공통 도우미 ──────────────────────────────────────────────────────────────
def read_text(path):
    """utf-8 → 일본어(cp932) → 한국어(cp949) 순서로 인코딩을 맞춰 읽는다."""
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "utf-8", "cp932", "cp949"):
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8(일부 깨짐)"


def fmt_time(seconds):
    if seconds is None:
        return "?"
    seconds = max(0, seconds)
    m, s = divmod(int(round(seconds)), 60)
    h, m = divmod(m, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"


def media_duration(path):
    """오디오/비디오 길이(초). 못 재면 None."""
    path = str(path)
    if mutagen is not None:
        try:
            mf = mutagen.File(path)
            if mf is not None and mf.info and getattr(mf.info, "length", 0):
                return float(mf.info.length)
        except Exception:
            pass
    if path.lower().endswith(".wav"):
        try:
            import wave
            with wave.open(path, "rb") as w:
                return w.getnframes() / float(w.getframerate())
        except Exception:
            pass
    # 마지막 수단: ffprobe가 설치되어 있으면 사용
    try:
        import subprocess
        out = subprocess.run(
            ["ffprobe", "-v", "error", "-show_entries", "format=duration",
             "-of", "default=noprint_wrappers=1:nokey=1", path],
            capture_output=True, text=True, timeout=30)
        return float(out.stdout.strip())
    except Exception:
        return None


def extract_number(name):
    """파일명에서 마지막 숫자 묶음을 뽑는다. scene_012.png → 12"""
    nums = re.findall(r"\d+", Path(name).stem)
    return int(nums[-1]) if nums else None


# ── 1) 프로젝트 폴더에서 파일 찾기 ──────────────────────────────────────────
def find_project_files(project_dir, report):
    sec = "① 파일 구성"
    p = Path(project_dir)
    files = {"script": None, "mapping": None, "storyboard": None,
             "subtitles": [], "images": [], "timeline": None,
             "audio": None, "video": None}

    all_files = [f for f in sorted(p.rglob("*")) if f.is_file()
                 and not any(part.startswith(".") for part in f.parts)]

    texts = [f for f in all_files if f.suffix.lower() in (".txt", ".md")]
    xlsxs = [f for f in all_files if f.suffix.lower() == ".xlsx"
             and not f.name.startswith("~$")]
    csvs = [f for f in all_files if f.suffix.lower() == ".csv"]

    # 씬 매핑: 이름에 scene-mapping / 씬매핑 등이 들어간 텍스트
    for f in texts:
        if re.search(r"scene[-_ ]?mapping|씬\s*매핑|씬매핑|マッピング", f.name, re.I):
            files["mapping"] = f
            break

    # 대본: 이름에 대본/台本/script 가 들어간 것 → 없으면 일본어 글자가 가장 많은 텍스트
    candidates = [f for f in texts if f != files["mapping"]]
    for f in candidates:
        if re.search(r"대본|台本|script|다이혼", f.name, re.I):
            files["script"] = f
            break
    if files["script"] is None and candidates:
        best, best_cnt = None, 0
        for f in candidates:
            try:
                text, _ = read_text(f)
            except Exception:
                continue
            cnt = len(JAPANESE_RE.findall(text))
            if cnt > best_cnt:
                best, best_cnt = f, cnt
        if best_cnt >= 100:
            files["script"] = best

    # 스토리보드(이미지 프롬프트 xlsx)와 타임라인
    for f in xlsxs:
        if re.search(r"prompt|storyboard|스토리보드|프롬프트", f.name, re.I):
            files["storyboard"] = f
        elif re.search(r"timeline|타임라인|timing|배치", f.name, re.I):
            files["timeline"] = f
    if files["storyboard"] is None and xlsxs:
        files["storyboard"] = xlsxs[0]
    for f in csvs:
        if files["timeline"] is None and re.search(
                r"timeline|타임라인|timing|배치|씬", f.name, re.I):
            files["timeline"] = f

    files["subtitles"] = [f for f in all_files if f.suffix.lower() in (".srt", ".ass")]
    files["images"] = [f for f in all_files if f.suffix.lower() in IMAGE_EXTS]
    audios = [f for f in all_files if f.suffix.lower() in AUDIO_EXTS]
    videos = [f for f in all_files if f.suffix.lower() in VIDEO_EXTS]
    files["audio"] = audios[0] if audios else None
    files["video"] = videos[0] if videos else None

    def rel(f):
        return str(f.relative_to(p)) if f else "없음"

    report.stat(sec, "대본", rel(files["script"]))
    report.stat(sec, "이미지", f"{len(files['images'])}장")
    report.stat(sec, "자막", ", ".join(rel(f) for f in files["subtitles"]) or "없음")
    report.stat(sec, "씬 매핑", rel(files["mapping"]))
    report.stat(sec, "스토리보드", rel(files["storyboard"]))
    report.stat(sec, "타임라인", rel(files["timeline"]))
    report.stat(sec, "오디오", rel(files["audio"]))
    report.stat(sec, "동영상", rel(files["video"]))

    if files["script"] is None:
        report.add(sec, "오류", "대본 파일을 찾지 못했습니다",
                   "파일명에 '대본'이 들어간 .txt 파일을 폴더에 넣어 주세요.")
    if not files["images"]:
        report.add(sec, "주의", "이미지 파일이 없습니다",
                   "png/jpg 이미지를 폴더(또는 images 하위 폴더)에 넣으면 이미지 검사를 합니다.")
    if not files["subtitles"]:
        report.add(sec, "주의", "자막 파일(.srt/.ass)이 없습니다",
                   "자막 타이밍·길이·위치 검사를 하려면 자막 파일을 넣어 주세요.")
    if files["script"] and files["images"] and files["subtitles"]:
        report.add(sec, "통과", "대본·이미지·자막이 모두 준비되어 있습니다")
    return files


# ── 2) 대본 검사 ─────────────────────────────────────────────────────────────
def split_sentences(text):
    parts = re.split(r"(?<=[。！？!?])\s*", text.replace("\n", ""))
    return [s.strip() for s in parts if s.strip()]


def check_script(path, report, cpm):
    sec = "② 대본"
    text, enc = read_text(path)
    lines = text.splitlines()

    chars_all = len(text.replace("\n", ""))
    chars_no_space = len(re.sub(r"\s", "", text))
    sentences = split_sentences(text)

    est_min = chars_no_space / cpm
    lo = chars_no_space / CPM_RANGE[1]
    hi = chars_no_space / CPM_RANGE[0]
    report.stat(sec, "인코딩", enc)
    report.stat(sec, "글자수(공백 포함)", f"{chars_all:,}자")
    report.stat(sec, "글자수(공백 제외)", f"{chars_no_space:,}자")
    report.stat(sec, "문장 수", f"{len(sentences):,}개")
    report.stat(sec, "예상 영상 길이",
                f"약 {est_min:.0f}분 (분당 {cpm}자 기준, {lo:.0f}~{hi:.0f}분 범위)")

    if "깨짐" in enc:
        report.add(sec, "오류", "대본 인코딩을 정확히 알 수 없어 일부 글자가 깨졌습니다",
                   "메모장에서 '다른 이름으로 저장' → 인코딩 UTF-8로 저장해 주세요.")

    # 한글 혼입 - TTS가 일본어 대본에서 한글을 만나면 사고가 납니다
    hangul_lines = [(i + 1, ln.strip()) for i, ln in enumerate(lines)
                    if HANGUL_RE.search(ln)]
    if hangul_lines:
        ex = "\n".join(f"{n}행: {ln[:60]}" for n, ln in hangul_lines[:10])
        more = f"\n... 외 {len(hangul_lines) - 10}곳" if len(hangul_lines) > 10 else ""
        report.add(sec, "오류", f"대본에 한글이 {len(hangul_lines)}곳 섞여 있습니다",
                   "TTS에 넣기 전에 반드시 지우거나 일본어로 바꿔야 합니다.\n" + ex + more)
    else:
        report.add(sec, "통과", "한글 혼입 없음 - 대본이 일본어로만 되어 있습니다")

    if not JAPANESE_RE.search(text):
        report.add(sec, "오류", "대본에서 일본어를 찾지 못했습니다",
                   "일본어 대본 파일이 맞는지 확인해 주세요.")

    # 마크다운 기호·지시문 잔재 - TTS가 그대로 읽어버립니다
    md_lines = [(i + 1, ln.strip()) for i, ln in enumerate(lines)
                if re.match(r"^\s*(#{1,6}\s|[-*]\s|\d+\.\s|>|```)", ln)
                or "**" in ln or re.match(r"^\s*-{3,}\s*$", ln)]
    if md_lines:
        ex = "\n".join(f"{n}행: {ln[:60]}" for n, ln in md_lines[:8])
        report.add(sec, "주의", f"마크다운 기호가 남은 줄이 {len(md_lines)}곳 있습니다",
                   "#, **, - 같은 기호를 TTS가 소리 내어 읽을 수 있습니다.\n" + ex)

    direction_lines = [(i + 1, ln.strip()) for i, ln in enumerate(lines)
                       if re.match(r"^\s*[\(（\[【].*[\)）\]】]\s*$", ln.strip())
                       and ln.strip()]
    if direction_lines:
        ex = "\n".join(f"{n}행: {ln[:60]}" for n, ln in direction_lines[:8])
        report.add(sec, "주의",
                   f"괄호로만 된 줄(지시문/헤더로 보임)이 {len(direction_lines)}곳 있습니다",
                   "(BGM), 【セクション】 같은 지시문은 TTS에 넣기 전에 지워야 합니다.\n" + ex)

    emojis = EMOJI_RE.findall(text)
    if emojis:
        report.add(sec, "주의", f"이모지·특수기호가 {len(emojis)}개 있습니다",
                   "예: " + " ".join(dict.fromkeys(emojis[:15])))

    if HALFWIDTH_KANA_RE.search(text):
        cnt = len(HALFWIDTH_KANA_RE.findall(text))
        report.add(sec, "주의", f"반각 가타카나가 {cnt}자 있습니다",
                   "반각 가타카나(ｱｲｳ...)는 TTS 발음이 어긋날 수 있으니 전각으로 바꿔 주세요.")

    # 너무 긴 문장 - TTS 호흡이 무너집니다
    long_sents = [s for s in sentences if len(s) > SENTENCE_MAX]
    if long_sents:
        ex = "\n".join(f"({len(s)}자) {s[:70]}..." for s in
                       sorted(long_sents, key=len, reverse=True)[:3])
        report.add(sec, "주의",
                   f"{SENTENCE_MAX}자를 넘는 긴 문장이 {len(long_sents)}개 있습니다",
                   "긴 문장은 TTS가 숨 쉴 틈 없이 읽어서 듣기 힘듭니다. 둘로 나눠 주세요.\n" + ex)
    else:
        report.add(sec, "통과", f"모든 문장이 {SENTENCE_MAX}자 이내입니다")

    # 문체 혼용 (です・ます체 vs 반말체)
    polite = plain = 0
    plain_examples = []
    for s in sentences:
        core = re.sub(r"[。！？!?」』\s]+$", "", s)
        if re.search(r"(です|ます|ました|でした|ません|ましょう|でしょう|ください|ですね|ますね|ですよ|ますよ|ですか|ますか)$", core):
            polite += 1
        elif re.search(r"(だ|である|だった|であった|だろう|する|した|いる|ある|ない|なる|なった|れる|られる|たい|う|く|い)$", core) and len(core) >= 8:
            plain += 1
            if len(plain_examples) < 5:
                plain_examples.append(s[:50])
    total_style = polite + plain
    if total_style >= 10:
        ratio = plain / total_style
        report.stat(sec, "문체", f"です・ます체 {polite}문장 / 반말체로 보이는 문장 {plain}문장")
        if ratio > 0.25:
            report.add(sec, "주의",
                       f"반말체(だ・である 등)로 끝나는 문장이 {ratio:.0%}입니다",
                       "です・ます체 기준 대본이라면 문체가 섞이지 않았는지 확인해 주세요.\n"
                       + "\n".join(plain_examples))
        else:
            report.add(sec, "통과", "문체가 です・ます체로 거의 통일되어 있습니다")

    # 중복 문장
    seen, dups = {}, []
    for s in sentences:
        if len(s) >= 15:
            seen[s] = seen.get(s, 0) + 1
    dups = [(s, c) for s, c in seen.items() if c >= 2]
    if dups:
        ex = "\n".join(f"({c}회) {s[:60]}" for s, c in
                       sorted(dups, key=lambda x: -x[1])[:5])
        report.add(sec, "주의", f"완전히 똑같은 문장이 {len(dups)}종류 반복됩니다",
                   "복사·붙여넣기 실수가 아닌지 확인해 주세요.\n" + ex)

    return text


# ── 3) 이미지 검사 ───────────────────────────────────────────────────────────
def check_images(image_files, report):
    sec = "③ 이미지"
    if not image_files:
        return
    report.stat(sec, "이미지 수", f"{len(image_files)}장")

    if Image is None:
        report.add(sec, "주의", "Pillow가 설치되지 않아 이미지 내용 검사를 건너뜁니다",
                   "install.bat 을 다시 실행해 주세요.")
        return

    broken, sizes, small_files, cmyk = [], {}, [], []
    for f in image_files:
        try:
            with Image.open(f) as im:
                im.verify()
            with Image.open(f) as im:
                sizes.setdefault(im.size, []).append(f.name)
                if im.mode == "CMYK":
                    cmyk.append(f.name)
        except Exception:
            broken.append(f.name)
            continue
        if f.stat().st_size < 30 * 1024:
            small_files.append(f.name)

    if broken:
        report.add(sec, "오류", f"열리지 않는(깨진) 이미지가 {len(broken)}장 있습니다",
                   "\n".join(broken[:10]))
    else:
        report.add(sec, "통과", "모든 이미지 파일이 정상적으로 열립니다")

    # 해상도·비율
    if sizes:
        main_size = max(sizes.items(), key=lambda kv: len(kv[1]))[0]
        report.stat(sec, "해상도", f"{main_size[0]}x{main_size[1]} 이 {len(sizes[main_size])}장 (가장 많음)")
        w, h = main_size
        if len(sizes) > 1:
            others = [f"{s[0]}x{s[1]} {len(v)}장 (예: {v[0]})"
                      for s, v in sizes.items() if s != main_size]
            report.add(sec, "주의", f"해상도가 서로 다른 이미지가 섞여 있습니다 ({len(sizes)}종류)",
                       "편집 시 확대·여백이 생길 수 있습니다.\n" + "\n".join(others[:8]))
        if w and h:
            ratio = w / h
            if abs(ratio - 16 / 9) > 0.02:
                report.add(sec, "주의",
                           f"이미지 비율이 16:9가 아닙니다 ({w}x{h}, 비율 {ratio:.2f})",
                           "유튜브 가로 영상 기준은 16:9(1920x1080)입니다. 세로 쇼츠라면 무시해도 됩니다.")
            if w < IMG_MIN_WIDTH:
                report.add(sec, "주의",
                           f"이미지 가로 해상도가 {w}px로 낮습니다",
                           f"최소 {IMG_MIN_WIDTH}px, 권장 {IMG_RECOMMEND[0]}x{IMG_RECOMMEND[1]} 이상을 추천합니다.")
            elif abs(ratio - 16 / 9) <= 0.02 and len(sizes) == 1:
                report.add(sec, "통과", f"모든 이미지가 {w}x{h} (16:9)로 통일되어 있습니다")

    if small_files:
        report.add(sec, "주의", f"용량이 30KB 미만인 이미지가 {len(small_files)}장 있습니다",
                   "생성이 실패했거나 거의 빈 이미지일 수 있습니다.\n" + "\n".join(small_files[:10]))
    if cmyk:
        report.add(sec, "주의", f"CMYK 색상 모드 이미지가 {len(cmyk)}장 있습니다",
                   "영상 편집 프로그램에서 색이 이상하게 나올 수 있습니다. RGB로 변환해 주세요.\n"
                   + "\n".join(cmyk[:5]))

    # 번호 연속성 검사
    numbered = [(extract_number(f.name), f.name) for f in image_files]
    numbered = [(n, name) for n, name in numbered if n is not None]
    if len(numbered) >= max(3, int(len(image_files) * 0.8)):
        nums = sorted(n for n, _ in numbered)
        dup_nums = sorted({n for n in nums if nums.count(n) > 1})
        missing = [n for n in range(nums[0], nums[-1] + 1) if n not in nums]
        report.stat(sec, "번호 범위", f"{nums[0]}번 ~ {nums[-1]}번")
        if missing:
            shown = ", ".join(str(n) for n in missing[:20])
            more = f" ... 외 {len(missing) - 20}개" if len(missing) > 20 else ""
            report.add(sec, "오류", f"번호가 빠진 이미지가 {len(missing)}장 있습니다",
                       f"빠진 번호: {shown}{more}")
        if dup_nums:
            report.add(sec, "주의", "같은 번호를 가진 이미지가 여러 장 있습니다",
                       "중복 번호: " + ", ".join(str(n) for n in dup_nums[:20]))
        if not missing and not dup_nums:
            report.add(sec, "통과", "이미지 번호가 빠짐없이 이어져 있습니다")


# ── 4) 자막 검사 (SRT / ASS) ────────────────────────────────────────────────
def parse_timecode(tc):
    m = re.match(r"(\d{1,2}):(\d{2}):(\d{2})[,.](\d{1,3})", tc.strip())
    if not m:
        return None
    h, mi, s, ms = m.groups()
    return int(h) * 3600 + int(mi) * 60 + int(s) + int(ms.ljust(3, "0")) / 1000.0


def parse_srt(text):
    """[(start, end, [줄, ...], 블록시작행번호), ...] 와 형식 오류 목록"""
    cues, errors = [], []
    blocks = re.split(r"\n\s*\n", text.strip())
    line_no = 1
    for block in blocks:
        block_lines = block.strip().splitlines()
        n_lines = block.count("\n") + 2
        start_line = line_no
        line_no += n_lines
        if not block_lines:
            continue
        idx = 0
        if re.match(r"^\d+$", block_lines[0].strip()):
            idx = 1
        if idx >= len(block_lines) or "-->" not in block_lines[idx]:
            errors.append(f"{start_line}행 부근: 시간 표시(-->)가 없는 블록")
            continue
        tc_parts = block_lines[idx].split("-->")
        start = parse_timecode(tc_parts[0])
        end = parse_timecode(tc_parts[1]) if len(tc_parts) > 1 else None
        if start is None or end is None:
            errors.append(f"{start_line}행 부근: 시간 형식이 잘못됨 ({block_lines[idx].strip()[:40]})")
            continue
        cues.append((start, end, block_lines[idx + 1:], start_line))
    return cues, errors


def parse_ass(text):
    """ASS 파일에서 (해상도, 스타일들, 이벤트들) 을 뽑는다."""
    play_res = [None, None]
    styles = {}
    events = []
    section = None
    style_format = event_format = None
    for raw in text.splitlines():
        line = raw.strip()
        if line.startswith("["):
            section = line.strip("[]").lower()
            continue
        if not line or line.startswith(";"):
            continue
        key, _, value = line.partition(":")
        key, value = key.strip().lower(), value.strip()
        if section == "script info":
            if key == "playresx":
                play_res[0] = int(float(value)) if value else None
            elif key == "playresy":
                play_res[1] = int(float(value)) if value else None
        elif section and "styles" in section:
            if key == "format":
                style_format = [v.strip().lower() for v in value.split(",")]
            elif key == "style" and style_format:
                vals = [v.strip() for v in value.split(",")]
                d = dict(zip(style_format, vals))
                styles[d.get("name", "?")] = d
        elif section == "events":
            if key == "format":
                event_format = [v.strip().lower() for v in value.split(",")]
            elif key == "dialogue" and event_format:
                vals = value.split(",", len(event_format) - 1)
                d = dict(zip(event_format, vals))
                events.append(d)
    return play_res, styles, events


def check_subtitle_common(cues, report, sec, label):
    """SRT/ASS 공통 타이밍·글자수 검사. cues = [(start, end, text한덩어리, 위치설명)]"""
    if not cues:
        report.add(sec, "오류", f"{label}: 자막을 한 개도 읽지 못했습니다")
        return
    report.stat(sec, f"{label} 자막 수", f"{len(cues)}개")
    report.stat(sec, f"{label} 자막 구간",
                f"{fmt_time(cues[0][0])} ~ {fmt_time(max(c[1] for c in cues))}")

    bad_order, overlaps, too_short, too_long_dur = [], [], [], []
    fast, very_fast, long_lines, many_lines, empty, hangul = [], [], [], [], [], []
    prev_end, prev_start = None, None
    for start, end, txt, where in cues:
        plain = re.sub(r"\{[^}]*\}", "", txt)          # ASS 태그 제거
        plain = re.sub(r"</?[a-zA-Z][^>]*>", "", plain)  # HTML 태그 제거
        cue_lines = [l for l in re.split(r"\\N|\n", plain)]
        joined = "".join(cue_lines)
        dur = end - start
        if end <= start:
            bad_order.append(f"{where}: 끝 시간이 시작보다 빠름 ({fmt_time(start)})")
        if prev_start is not None and start < prev_start - 0.001:
            bad_order.append(f"{where}: 앞 자막보다 시작 시간이 빠름 ({fmt_time(start)})")
        if prev_end is not None and start < prev_end - 0.05:
            overlaps.append(f"{where}: 앞 자막과 {prev_end - start:.2f}초 겹침 ({fmt_time(start)})")
        prev_end, prev_start = end, start
        if not joined.strip():
            empty.append(f"{where} ({fmt_time(start)})")
            continue
        if dur < SUB_MIN_DUR:
            too_short.append(f"{where}: {dur:.2f}초 ({joined[:20]})")
        if dur > SUB_MAX_DUR:
            too_long_dur.append(f"{where}: {dur:.1f}초 ({joined[:20]})")
        n_chars = len(re.sub(r"\s", "", joined))
        if dur > 0:
            cps = n_chars / dur
            if cps > SUB_CPS_BAD:
                very_fast.append(f"{where}: 초당 {cps:.1f}자 ({joined[:25]})")
            elif cps > SUB_CPS_WARN:
                fast.append(f"{where}: 초당 {cps:.1f}자 ({joined[:25]})")
        for l in cue_lines:
            if len(l) > SUB_LINE_MAX:
                long_lines.append(f"{where}: {len(l)}자 ({l[:30]}...)")
                break
        if len([l for l in cue_lines if l.strip()]) > 2:
            many_lines.append(f"{where} ({fmt_time(start)})")
        if HANGUL_RE.search(joined):
            hangul.append(f"{where}: {joined[:40]}")

    def add_list(severity, title, items, tip=""):
        if items:
            ex = "\n".join(items[:10])
            more = f"\n... 외 {len(items) - 10}건" if len(items) > 10 else ""
            report.add(sec, severity, f"{label}: {title} ({len(items)}건)",
                       (tip + "\n" if tip else "") + ex + more)

    add_list("오류", "시간 순서가 잘못된 자막", bad_order)
    add_list("오류", "앞 자막과 시간이 겹치는 자막", overlaps,
             "두 자막이 동시에 화면에 떠서 겹쳐 보입니다.")
    add_list("오류", "내용이 빈 자막", empty)
    add_list("오류", "한글이 섞인 자막", hangul,
             "일본어 자막에 한글이 남아 있습니다.")
    add_list("주의", f"표시 시간이 {SUB_MIN_DUR}초 미만인 자막", too_short,
             "너무 짧아서 읽기 전에 사라집니다.")
    add_list("주의", f"표시 시간이 {SUB_MAX_DUR}초를 넘는 자막", too_long_dur)
    add_list("주의", f"너무 빠른 자막 (초당 {SUB_CPS_BAD}자 초과)", very_fast,
             "이 속도는 사실상 읽을 수 없습니다. 자막을 나누거나 표시 시간을 늘려 주세요.")
    add_list("정보", f"조금 빠른 자막 (초당 {SUB_CPS_WARN}자 초과)", fast)
    add_list("주의", f"한 줄이 {SUB_LINE_MAX}자를 넘는 자막", long_lines,
             "일본어 자막은 한 줄 20자 이내가 읽기 편합니다.")
    add_list("주의", "3줄 이상인 자막", many_lines,
             "자막은 최대 2줄을 권장합니다.")

    problem_count = sum(len(x) for x in
                        (bad_order, overlaps, empty, hangul, very_fast))
    if problem_count == 0:
        report.add(sec, "통과", f"{label}: 타이밍·내용에 큰 문제가 없습니다")
    return max(c[1] for c in cues)


def check_subtitles(sub_files, report):
    sec = "④ 자막"
    last_end = None
    for f in sub_files:
        text, enc = read_text(f)
        label = f.name
        if f.suffix.lower() == ".srt":
            cues, errors = parse_srt(text)
            for e in errors[:10]:
                report.add(sec, "오류", f"{label}: 형식 오류", e)
            simple = [(s, e2, "\n".join(t), f"{i + 1}번 자막")
                      for i, (s, e2, t, _) in enumerate(cues)]
            end = check_subtitle_common(simple, report, sec, label)
            if end:
                last_end = max(last_end or 0, end)
        elif f.suffix.lower() == ".ass":
            end = check_ass_file(f, text, report, sec)
            if end:
                last_end = max(last_end or 0, end)
    return last_end


def check_ass_file(f, text, report, sec):
    label = f.name
    play_res, styles, events = parse_ass(text)
    W = play_res[0] or 1920
    H = play_res[1] or 1080
    report.stat(sec, f"{label} 기준 해상도", f"{W}x{H}"
                + ("" if play_res[0] else " (파일에 명시 안 됨 → 1920x1080으로 가정)"))

    # 스타일(기본 자막 위치) 검사
    used_styles = {e.get("style", "").lstrip("*") for e in events}
    for name in sorted(used_styles):
        st = styles.get(name)
        if st is None:
            report.add(sec, "주의", f"{label}: 정의되지 않은 스타일 '{name}'을 쓰고 있습니다",
                       "편집 프로그램 기본값으로 표시되어 의도한 위치·모양과 달라질 수 있습니다.")
            continue
        try:
            align = int(float(st.get("alignment", 2)))
        except ValueError:
            align = 2
        try:
            margin_v = int(float(st.get("marginv", 0)))
        except ValueError:
            margin_v = 0
        if align not in (1, 2, 3):
            pos_name = {4: "왼쪽 중간", 5: "정중앙", 6: "오른쪽 중간",
                        7: "왼쪽 위", 8: "위쪽 가운데", 9: "오른쪽 위"}.get(align, f"Alignment={align}")
            report.add(sec, "주의",
                       f"{label}: 스타일 '{name}'의 자막 위치가 화면 하단이 아닙니다 ({pos_name})",
                       "일반 자막은 하단 가운데(Alignment=2)가 기본입니다. 의도한 배치인지 확인해 주세요.")
        if align in (1, 2, 3) and margin_v < 10:
            report.add(sec, "주의",
                       f"{label}: 스타일 '{name}'의 아래 여백(MarginV)이 {margin_v}px입니다",
                       "화면 맨 끝에 붙어 TV·모바일에서 잘릴 수 있습니다. 40px 이상을 권장합니다.")

    # \pos() 로 위치를 직접 지정한 자막 검사 - "이미지별 자막 위치"가 화면을 벗어나는지
    out_of_screen, upper_area = [], []
    cues = []
    for i, e in enumerate(events):
        start = parse_timecode(e.get("start", ""))
        end = parse_timecode(e.get("end", ""))
        raw_text = e.get("text", "")
        if start is None or end is None:
            report.add(sec, "오류", f"{label}: {i + 1}번 자막의 시간 형식이 잘못되었습니다",
                       raw_text[:50])
            continue
        where = f"{i + 1}번 자막"
        cues.append((start, end, raw_text, where))
        for m in re.finditer(r"\\pos\(\s*(-?[\d.]+)\s*,\s*(-?[\d.]+)\s*\)", raw_text):
            x, y = float(m.group(1)), float(m.group(2))
            if not (0 <= x <= W and 0 <= y <= H):
                out_of_screen.append(
                    f"{where} ({fmt_time(start)}): 위치 ({x:.0f},{y:.0f}) - 화면({W}x{H}) 밖")
            elif y < H * 0.70:
                upper_area.append(
                    f"{where} ({fmt_time(start)}): 위치 ({x:.0f},{y:.0f}) - 화면 위쪽 {y / H:.0%} 지점")
    if out_of_screen:
        report.add(sec, "오류",
                   f"{label}: 화면 밖에 찍힌 자막이 {len(out_of_screen)}개 있습니다",
                   "\n".join(out_of_screen[:10]))
    if upper_area:
        report.add(sec, "주의",
                   f"{label}: 화면 위쪽(상단 70% 영역)에 배치된 자막이 {len(upper_area)}개 있습니다",
                   "이미지의 중요한 부분을 가릴 수 있습니다. 의도한 배치인지 확인해 주세요.\n"
                   + "\n".join(upper_area[:10]))
    if not out_of_screen and not upper_area and events:
        report.add(sec, "통과", f"{label}: 자막 위치가 모두 화면 안 하단 영역에 있습니다")

    return check_subtitle_common(cues, report, sec, label)


# ── 5) 씬 매핑·스토리보드 검사 ──────────────────────────────────────────────
def normalize_for_match(s):
    s = unicodedata.normalize("NFKC", s)
    return re.sub(r"[\s、。「」『』・,.!?！?？…―ー\-()（）\"']", "", s)


def check_scene_mapping(path, script_text, image_count, report):
    sec = "⑤ 씬 구성"
    text, _ = read_text(path)
    entries = []
    for ln in text.splitlines():
        m = re.match(r"^\s*(?:씬|シーン|scene)?\s*0*(\d+)\s*[.:：)\]]\s*(.+)$",
                     ln.strip(), re.I)
        if m:
            entries.append((int(m.group(1)), m.group(2).strip()))
    if not entries:
        report.add(sec, "주의", "씬 매핑 파일에서 '씬 번호: 시작 문구' 줄을 찾지 못했습니다",
                   "예: 「씬 1: 皆さん、こんにちは」 형식이어야 합니다.")
        return

    report.stat(sec, "씬 매핑 항목", f"{len(entries)}개")
    nums = [n for n, _ in entries]
    expected = list(range(min(nums), min(nums) + len(nums)))
    if nums != expected:
        report.add(sec, "주의", "씬 번호가 순서대로 이어지지 않습니다",
                   f"현재 번호: {nums[:30]}")

    if image_count and len(entries) != image_count:
        report.add(sec, "오류",
                   f"씬 수({len(entries)}개)와 이미지 수({image_count}장)가 다릅니다",
                   "씬마다 이미지가 1장씩이라면 두 숫자가 같아야 합니다.")
    elif image_count:
        report.add(sec, "통과", f"씬 수와 이미지 수가 {image_count}개로 일치합니다")

    if script_text:
        norm_script = normalize_for_match(script_text)
        not_found, out_of_order = [], []
        prev_pos = -1
        for n, phrase in entries:
            key = normalize_for_match(phrase)[:20]
            if len(key) < 4:
                continue
            pos = norm_script.find(key, max(prev_pos, 0))
            if pos < 0:
                pos_any = norm_script.find(key)
                if pos_any < 0:
                    not_found.append(f"씬 {n}: {phrase[:40]}")
                else:
                    out_of_order.append(f"씬 {n}: {phrase[:40]}")
                    prev_pos = pos_any
            else:
                prev_pos = pos
        if not_found:
            report.add(sec, "오류",
                       f"대본에서 찾을 수 없는 씬 시작 문구가 {len(not_found)}개 있습니다",
                       "대본을 수정한 뒤 씬 매핑을 갱신하지 않았을 수 있습니다.\n"
                       + "\n".join(not_found[:10]))
        if out_of_order:
            report.add(sec, "주의",
                       f"대본 순서와 어긋나는 씬 시작 문구가 {len(out_of_order)}개 있습니다",
                       "\n".join(out_of_order[:10]))
        if not not_found and not out_of_order:
            report.add(sec, "통과", "모든 씬 시작 문구가 대본에 순서대로 존재합니다")


def check_storyboard(path, image_count, report):
    sec = "⑤ 씬 구성"
    if openpyxl is None:
        report.add(sec, "주의", "openpyxl이 설치되지 않아 스토리보드(xlsx) 검사를 건너뜁니다",
                   "install.bat 을 다시 실행해 주세요.")
        return
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        report.add(sec, "오류", f"스토리보드 파일을 열 수 없습니다: {path.name}", str(e))
        return
    ws = None
    for name in wb.sheetnames:
        if name.lower() in ("storyboard", "스토리보드"):
            ws = wb[name]
            break
    if ws is None:
        ws = wb[wb.sheetnames[0]]

    rows = [[("" if c is None else str(c).strip()) for c in row]
            for row in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(r)]
    if not rows:
        report.add(sec, "오류", f"스토리보드가 비어 있습니다: {path.name}")
        return
    # 첫 줄이 헤더면 건너뜀
    if rows and not re.match(r"^\d+$", rows[0][0]):
        rows = rows[1:]

    report.stat(sec, "스토리보드 씬 수", f"{len(rows)}개 ({path.name})")
    empty_prompts = [r[0] for r in rows if len(r) < 3 or not r[2]]
    if empty_prompts:
        report.add(sec, "오류",
                   f"이미지 프롬프트가 비어 있는 씬이 {len(empty_prompts)}개 있습니다",
                   "씬 번호: " + ", ".join(empty_prompts[:20]))
    prompts = [r[2] for r in rows if len(r) >= 3 and r[2]]
    dup = {p for p in prompts if prompts.count(p) > 1}
    if dup:
        report.add(sec, "주의", f"완전히 똑같은 이미지 프롬프트가 {len(dup)}종류 있습니다",
                   "같은 그림이 여러 번 생성됩니다. 의도한 것인지 확인해 주세요.\n"
                   + "\n".join(p[:80] for p in list(dup)[:3]))
    if image_count and len(rows) != image_count:
        report.add(sec, "주의",
                   f"스토리보드 씬 수({len(rows)}개)와 실제 이미지 수({image_count}장)가 다릅니다",
                   "아직 생성하지 않은 이미지가 있거나, 여분 이미지가 섞여 있을 수 있습니다.")
    if not empty_prompts and (not image_count or len(rows) == image_count):
        report.add(sec, "통과", "스토리보드 프롬프트가 모두 채워져 있습니다")


# ── 6) 타임라인(동영상 배치) 검사 ───────────────────────────────────────────
def load_timeline_rows(path):
    """CSV/XLSX에서 [{컬럼: 값}] 형태로 읽는다."""
    if path.suffix.lower() == ".csv":
        text, _ = read_text(path)
        rows = list(csv.reader(text.splitlines()))
    else:
        if openpyxl is None:
            return None, "openpyxl 미설치"
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [[("" if c is None else str(c)) for c in r]
                for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if len(rows) < 2:
        return None, "내용 부족"
    header = [str(c).strip().lower() for c in rows[0]]

    def find_col(*keys):
        for i, h in enumerate(header):
            if any(k in h for k in keys):
                return i
        return None

    cols = {
        "num": find_col("씬", "번호", "no", "scene", "順"),
        "file": find_col("파일", "이미지", "file", "image", "영상", "clip", "素材"),
        "start": find_col("시작", "start", "開始"),
        "dur": find_col("길이", "duration", "초", "尺", "時間"),
    }
    out = []
    for r in rows[1:]:
        item = {}
        for k, i in cols.items():
            item[k] = str(r[i]).strip() if i is not None and i < len(r) else ""
        out.append(item)
    return out, None


def parse_seconds(v):
    """'75', '1:15', '0:01:15', '75.5' → 초"""
    v = str(v).strip()
    if not v:
        return None
    if re.match(r"^\d+(\.\d+)?$", v):
        return float(v)
    parts = v.split(":")
    try:
        parts = [float(p) for p in parts]
    except ValueError:
        return None
    sec = 0.0
    for p in parts:
        sec = sec * 60 + p
    return sec


def check_timeline(path, image_files, audio_dur, report):
    sec = "⑥ 타임라인·길이"
    rows, err = load_timeline_rows(path)
    if rows is None:
        report.add(sec, "주의", f"타임라인 파일을 해석하지 못했습니다: {path.name}", err or "")
        return
    report.stat(sec, "타임라인 항목", f"{len(rows)}개 ({path.name})")

    image_names = {f.name for f in image_files}
    image_stems = {f.stem for f in image_files}
    missing_files = []
    for i, r in enumerate(rows, start=2):
        fn = r.get("file", "")
        if fn and Path(fn).suffix.lower() in IMAGE_EXTS | VIDEO_EXTS:
            base = Path(fn).name
            if base not in image_names and Path(base).stem not in image_stems:
                missing_files.append(f"{i}행: {fn}")
    if missing_files:
        report.add(sec, "오류",
                   f"타임라인에 적혀 있지만 폴더에 없는 파일이 {len(missing_files)}개 있습니다",
                   "\n".join(missing_files[:10]))

    starts = [(i + 2, parse_seconds(r.get("start"))) for i, r in enumerate(rows)]
    starts_valid = [(ln, s) for ln, s in starts if s is not None]
    if len(starts_valid) >= 2:
        disorder = [f"{ln}행: {fmt_time(s)}"
                    for (ln, s), (_, prev) in zip(starts_valid[1:], starts_valid[:-1])
                    if s < prev]
        if disorder:
            report.add(sec, "오류", f"시작 시간이 앞 항목보다 빠른 줄이 {len(disorder)}개 있습니다",
                       "\n".join(disorder[:10]))
        else:
            report.add(sec, "통과", "타임라인 시작 시간이 순서대로 늘어납니다")

        durs = [parse_seconds(r.get("dur")) for r in rows]
        if all(s is not None for _, s in starts_valid) and any(durs):
            gaps, overlaps = [], []
            for idx in range(len(starts_valid) - 1):
                ln, s = starts_valid[idx]
                d = durs[idx] if idx < len(durs) else None
                if d is None:
                    continue
                nxt = starts_valid[idx + 1][1]
                end = s + d
                if end - nxt > 0.5:
                    overlaps.append(f"{ln}행: 다음 씬과 {end - nxt:.1f}초 겹침")
                elif nxt - end > 0.5:
                    gaps.append(f"{ln}행: 다음 씬까지 {nxt - end:.1f}초 빈 구간")
            if overlaps:
                report.add(sec, "주의", f"씬끼리 겹치는 구간이 {len(overlaps)}곳 있습니다",
                           "\n".join(overlaps[:10]))
            if gaps:
                report.add(sec, "주의", f"아무것도 없는 빈 구간이 {len(gaps)}곳 있습니다",
                           "\n".join(gaps[:10]))

        total_end = None
        last_ln, last_s = starts_valid[-1]
        last_d = parse_seconds(rows[-1].get("dur"))
        if last_d is not None:
            total_end = last_s + last_d
            report.stat(sec, "타임라인 총 길이", fmt_time(total_end))
        if total_end and audio_dur:
            diff = total_end - audio_dur
            if abs(diff) > 3:
                report.add(sec, "주의",
                           f"타임라인 총 길이({fmt_time(total_end)})와 "
                           f"오디오 길이({fmt_time(audio_dur)})가 {abs(diff):.0f}초 다릅니다",
                           "이미지가 끝났는데 소리만 나오거나, 소리 없이 화면만 남을 수 있습니다.")
            else:
                report.add(sec, "통과", "타임라인 총 길이와 오디오 길이가 맞습니다")


# ── 7) 전체 길이 정합 검사 ──────────────────────────────────────────────────
def check_durations(files, script_text, sub_last_end, report, cpm):
    sec = "⑥ 타임라인·길이"
    audio_dur = video_dur = None
    if files["audio"]:
        audio_dur = media_duration(files["audio"])
        report.stat(sec, "오디오 길이",
                    fmt_time(audio_dur) if audio_dur else "측정 실패")
        if audio_dur is None:
            report.add(sec, "정보", "오디오 길이를 잴 수 없어 길이 비교를 건너뜁니다",
                       "install.bat 을 다시 실행하면 대부분 해결됩니다. (mutagen 설치)")
    if files["video"]:
        video_dur = media_duration(files["video"])
        report.stat(sec, "동영상 길이",
                    fmt_time(video_dur) if video_dur else "측정 실패")

    if script_text and audio_dur:
        chars = len(re.sub(r"\s", "", script_text))
        est = chars / cpm * 60
        diff_ratio = abs(est - audio_dur) / max(audio_dur, 1)
        if diff_ratio > AUDIO_DIFF_RATIO:
            report.add(sec, "주의",
                       f"대본 예상 길이(약 {fmt_time(est)})와 오디오 길이({fmt_time(audio_dur)})가 "
                       f"{diff_ratio:.0%} 차이 납니다",
                       "대본 일부만 녹음됐거나, 다른 버전의 대본·오디오가 섞였을 수 있습니다.\n"
                       f"(분당 {cpm}자 기준 계산이므로 TTS 속도 설정에 따라 어느 정도 차이는 정상입니다)")
        else:
            report.add(sec, "통과", "대본 분량과 오디오 길이가 자연스럽게 맞습니다")

    base_dur = audio_dur or video_dur
    if sub_last_end and base_dur:
        src = "오디오" if audio_dur else "동영상"
        if sub_last_end > base_dur + 1.5:
            report.add(sec, "오류",
                       f"자막이 {src}보다 {sub_last_end - base_dur:.0f}초 더 깁니다",
                       f"마지막 자막 끝: {fmt_time(sub_last_end)} / {src} 길이: {fmt_time(base_dur)}\n"
                       "영상이 끝난 뒤에도 자막 타이밍이 남아 있습니다.")
        elif base_dur - sub_last_end > base_dur * 0.10 + 5:
            report.add(sec, "주의",
                       f"마지막 자막({fmt_time(sub_last_end)}) 이후 {src} 끝({fmt_time(base_dur)})까지 "
                       f"{base_dur - sub_last_end:.0f}초 동안 자막이 없습니다",
                       "뒷부분 자막이 빠지지 않았는지 확인해 주세요.")
        else:
            report.add(sec, "통과", f"자막 구간이 {src} 길이와 잘 맞습니다")

    if audio_dur and video_dur and abs(audio_dur - video_dur) > 3:
        report.add(sec, "주의",
                   f"오디오({fmt_time(audio_dur)})와 동영상({fmt_time(video_dur)}) 길이가 "
                   f"{abs(audio_dur - video_dur):.0f}초 다릅니다")
    return audio_dur


# ── 8) HTML 리포트 만들기 ───────────────────────────────────────────────────
SEVERITY_COLORS = {"오류": "#d93025", "주의": "#e37400", "정보": "#1a73e8", "통과": "#188038"}
SEVERITY_BG = {"오류": "#fce8e6", "주의": "#fef7e0", "정보": "#e8f0fe", "통과": "#e6f4ea"}


def render_html(report, project_dir, out_path):
    esc = html_mod.escape
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    counts = {s: report.count(s) for s in Report.SEVERITIES}

    if counts["오류"]:
        verdict = ("❌ 수정이 필요합니다", "#d93025",
                   f"오류 {counts['오류']}건을 고친 뒤 다시 검토해 주세요.")
    elif counts["주의"]:
        verdict = ("⚠️ 거의 다 됐습니다", "#e37400",
                   f"주의 {counts['주의']}건만 확인하면 업로드 준비 완료입니다.")
    else:
        verdict = ("✅ 통과", "#188038", "발견된 문제가 없습니다. 업로드 준비 완료!")

    cards = "".join(
        f'<div class="card" style="border-top:4px solid {SEVERITY_COLORS[s]}">'
        f'<div class="num" style="color:{SEVERITY_COLORS[s]}">{counts[s]}</div>'
        f'<div class="lbl">{s}</div></div>'
        for s in Report.SEVERITIES)

    body_sections = []
    for section in report.sections:
        stat_rows = "".join(
            f"<tr><th>{esc(str(k))}</th><td>{esc(str(v))}</td></tr>"
            for k, v in report.stats.get(section, []))
        stat_html = f'<table class="stats">{stat_rows}</table>' if stat_rows else ""

        items = [f for f in report.findings if f[0] == section]
        order = {s: i for i, s in enumerate(Report.SEVERITIES)}
        items.sort(key=lambda f: order[f[1]])
        rows = []
        for _, sev, title, detail in items:
            detail_html = (f'<div class="detail">{esc(detail)}</div>'
                           if detail else "")
            rows.append(
                f'<div class="finding" style="background:{SEVERITY_BG[sev]}">'
                f'<span class="badge" style="background:{SEVERITY_COLORS[sev]}">{sev}</span>'
                f'<div class="msg"><div class="title">{esc(title)}</div>{detail_html}</div></div>')
        body_sections.append(
            f'<section><h2>{esc(section)}</h2>{stat_html}{"".join(rows) or "<p class=empty>검사 항목 없음</p>"}</section>')

    html_out = f"""<!DOCTYPE html>
<html lang="ko"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>일본 영상 검토 리포트</title>
<style>
  body {{ font-family:'Malgun Gothic','Apple SD Gothic Neo',sans-serif; margin:0;
         background:#f6f7f9; color:#222; }}
  .wrap {{ max-width:900px; margin:0 auto; padding:24px 16px 60px; }}
  header {{ background:#fff; border-radius:12px; padding:24px; box-shadow:0 1px 4px rgba(0,0,0,.08); }}
  h1 {{ margin:0 0 6px; font-size:22px; }}
  .meta {{ color:#666; font-size:13px; }}
  .verdict {{ margin-top:14px; font-size:18px; font-weight:bold; }}
  .verdict small {{ display:block; font-weight:normal; font-size:13px; color:#555; margin-top:4px; }}
  .cards {{ display:flex; gap:12px; margin:18px 0 6px; flex-wrap:wrap; }}
  .card {{ flex:1; min-width:110px; background:#fff; border-radius:10px; text-align:center;
           padding:14px 8px; box-shadow:0 1px 4px rgba(0,0,0,.08); }}
  .card .num {{ font-size:28px; font-weight:bold; }}
  .card .lbl {{ color:#666; font-size:13px; margin-top:2px; }}
  section {{ background:#fff; border-radius:12px; padding:20px 24px; margin-top:18px;
             box-shadow:0 1px 4px rgba(0,0,0,.08); }}
  h2 {{ font-size:17px; margin:0 0 12px; border-bottom:2px solid #eee; padding-bottom:8px; }}
  table.stats {{ border-collapse:collapse; margin-bottom:14px; font-size:13px; width:100%; }}
  table.stats th {{ text-align:left; color:#666; font-weight:normal; padding:3px 14px 3px 0;
                    white-space:nowrap; vertical-align:top; width:1%; }}
  table.stats td {{ padding:3px 0; word-break:break-all; }}
  .finding {{ display:flex; gap:10px; border-radius:8px; padding:10px 12px; margin:8px 0;
              align-items:flex-start; }}
  .badge {{ color:#fff; font-size:12px; padding:2px 8px; border-radius:10px; white-space:nowrap;
            margin-top:1px; }}
  .title {{ font-weight:bold; font-size:14px; }}
  .detail {{ white-space:pre-wrap; font-size:12.5px; color:#444; margin-top:4px;
             font-family:Consolas,'Malgun Gothic',monospace; }}
  .empty {{ color:#999; font-size:13px; }}
</style></head><body><div class="wrap">
<header>
  <h1>🎬 일본 영상 프로젝트 검토 리포트</h1>
  <div class="meta">폴더: {esc(str(project_dir))}<br>검토 시각: {now}</div>
  <div class="verdict" style="color:{verdict[1]}">{verdict[0]}<small>{verdict[2]}</small></div>
</header>
<div class="cards">{cards}</div>
{"".join(body_sections)}
</div></body></html>"""
    Path(out_path).write_text(html_out, encoding="utf-8")


# ── 메인 ─────────────────────────────────────────────────────────────────────
def run_review(project_dir, cpm=DEFAULT_CPM, out=None, auto_open=True):
    project_dir = Path(project_dir).expanduser()
    if not project_dir.is_dir():
        print(f"폴더를 찾을 수 없습니다: {project_dir}")
        return 2

    print("=" * 52)
    print(" 일본 영상 프로젝트 검토를 시작합니다")
    print(f" 대상 폴더: {project_dir}")
    print("=" * 52)

    report = Report()
    files = find_project_files(project_dir, report)

    script_text = None
    if files["script"]:
        print(f"· 대본 검사 중: {files['script'].name}")
        script_text = check_script(files["script"], report, cpm)

    if files["images"]:
        print(f"· 이미지 검사 중: {len(files['images'])}장")
        check_images(files["images"], report)

    sub_last_end = None
    if files["subtitles"]:
        print(f"· 자막 검사 중: {len(files['subtitles'])}개 파일")
        sub_last_end = check_subtitles(files["subtitles"], report)

    if files["mapping"]:
        print(f"· 씬 매핑 검사 중: {files['mapping'].name}")
        check_scene_mapping(files["mapping"], script_text,
                            len(files["images"]), report)
    if files["storyboard"]:
        print(f"· 스토리보드 검사 중: {files['storyboard'].name}")
        check_storyboard(files["storyboard"], len(files["images"]), report)

    print("· 길이 정합 검사 중")
    audio_dur = check_durations(files, script_text, sub_last_end, report, cpm)

    if files["timeline"]:
        print(f"· 타임라인 검사 중: {files['timeline'].name}")
        check_timeline(files["timeline"], files["images"], audio_dur, report)

    out_path = Path(out) if out else project_dir / "검토리포트.html"
    render_html(report, project_dir, out_path)

    print()
    print("-" * 52)
    for s in Report.SEVERITIES:
        print(f"  {s}: {report.count(s)}건")
    print("-" * 52)
    for section, sev, title, _ in report.findings:
        if sev == "오류":
            print(f"  [오류] {title}")
    print()
    print(f"자세한 내용: {out_path}")

    if auto_open:
        try:
            if os.name == "nt":
                os.startfile(out_path)  # noqa
        except Exception:
            pass
    return 1 if report.count("오류") else 0


def main():
    parser = argparse.ArgumentParser(description="일본 영상 프로젝트 검토 프로그램")
    parser.add_argument("folder", nargs="?", help="검토할 프로젝트 폴더")
    parser.add_argument("--cpm", type=int, default=DEFAULT_CPM,
                        help=f"TTS 분당 글자수 기준 (기본 {DEFAULT_CPM})")
    parser.add_argument("--out", help="리포트 저장 경로 (기본: 폴더 안 검토리포트.html)")
    parser.add_argument("--no-open", action="store_true",
                        help="검토 후 리포트를 자동으로 열지 않음")
    args = parser.parse_args()

    folder = args.folder
    if not folder:
        folder = input("검토할 프로젝트 폴더 경로를 붙여넣고 Enter: ").strip().strip('"')
    sys.exit(run_review(folder, cpm=args.cpm, out=args.out,
                        auto_open=not args.no_open))


if __name__ == "__main__":
    main()
